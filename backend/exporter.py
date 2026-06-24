import openpyxl
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
from openpyxl.utils import get_column_letter
import io
import re
from models import slot_to_day_period, SubjectType

DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
PERIODS = ["9:00-10:00", "10:00-11:00", "11:00-12:00", "12:00-1:00", "1:00-2:00", "2:00-3:00", "3:00-4:00", "4:00-5:00", "5:00-6:00"]

# Colors
THEORY_FILL = PatternFill(start_color="D9F2FF", end_color="D9F2FF", fill_type="solid")
PRAC_FILL = PatternFill(start_color="D9F7DC", end_color="D9F7DC", fill_type="solid")
TUT_FILL = PatternFill(start_color="FFF0D4", end_color="FFF0D4", fill_type="solid")
LUNCH_FILL = PatternFill(start_color="FFD6D6", end_color="FFD6D6", fill_type="solid")

def extract_batch_idx(batch_str: str) -> int:
    m = re.search(r'\d+', batch_str)
    if m:
        return int(m.group()) - 1
    return 0

def export_timetable_to_excel(data, timetable_index):
    wb = openpyxl.Workbook()
    wb.remove(wb.active) # Remove default sheet
    
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    header_fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
    
    # ── DIVISION SHEETS ───────────────────────────────────────
    for div_key, div in data.divisions.items():
        ws = wb.create_sheet(title=div_key[:31]) # Sheet names max 31 chars
        
        # Determine number of batches
        batches = [b.batch_id for b in div.batches]
        if not batches:
            batches = ["ALL"]
        
        num_batches = len(batches) if batches != ["ALL"] else 1
        
        # Build headers
        ws.cell(row=1, column=1, value="Time/Day").border = thin_border
        ws.cell(row=1, column=1).fill = header_fill
        ws.cell(row=1, column=1).font = Font(bold=True)
        
        for i, day in enumerate(DAYS):
            start_col = 2 + (i * num_batches)
            end_col = start_col + num_batches - 1
            if start_col < end_col:
                ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
            cell = ws.cell(row=1, column=start_col, value=day)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = Font(bold=True)
            cell.fill = header_fill
            for c in range(start_col, end_col + 1):
                ws.cell(row=1, column=c).border = thin_border
                
        # Setup periods
        for p_idx, p_label in enumerate(PERIODS):
            row_idx = p_idx + 2
            cell = ws.cell(row=row_idx, column=1, value=p_label)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
            
            # Draw borders for all cells
            for c in range(2, 2 + 5 * num_batches):
                ws.cell(row=row_idx, column=c).border = thin_border
                
        # Lunch
        lunch_row = 6 # 1-2 PM is period 5, which is row 6
        if num_batches > 1:
            ws.merge_cells(start_row=lunch_row, start_column=2, end_row=lunch_row, end_column=1 + 5 * num_batches)
        lunch_cell = ws.cell(row=lunch_row, column=2, value="LUNCH BREAK")
        lunch_cell.alignment = Alignment(horizontal="center", vertical="center")
        lunch_cell.font = Font(bold=True)
        lunch_cell.fill = LUNCH_FILL
        
        # Get records
        week_records = timetable_index.get_division_week(div_key)
        
        for slot_id, rec_list in week_records.items():
            day_str, period = slot_to_day_period(slot_id)
            day_idx = slot_id // 9
            if period == 5: continue # Skip lunch just in case
            
            row_idx = period + 1 # +1 because period is 1-indexed, but row 1 is header
            
            day_start_col = 2 + (day_idx * num_batches)
            
            for rec in rec_list:
                # Only write cells on the START slot of the session
                if slot_id != rec.slot_id:
                    continue
                    
                if rec.subject_type == SubjectType.THEORY:
                    # Spans all batches
                    if num_batches > 1:
                        try:
                            ws.merge_cells(start_row=row_idx, start_column=day_start_col, end_row=row_idx, end_column=day_start_col + num_batches - 1)
                        except ValueError:
                            pass
                    cell = ws.cell(row=row_idx, column=day_start_col)
                    text = f"{rec.course_code}\n{' '.join(rec.room_nos)}\n{' '.join(rec.fac_ids)}"
                    cell.value = text
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    cell.fill = THEORY_FILL
                    
                elif rec.subject_type in (SubjectType.PRACTICAL, SubjectType.TUTORIAL):
                    # Finds its specific batch column based on batch number
                    for batch in rec.batch_ids:
                        bIdx = extract_batch_idx(batch) if batch != "ALL" else 0
                        bIdx = max(0, min(num_batches - 1, bIdx))
                        
                        col_idx = day_start_col + bIdx
                        end_row_idx = row_idx + rec.duration - 1
                        
                        if rec.duration > 1:
                            try:
                                ws.merge_cells(start_row=row_idx, start_column=col_idx, end_row=end_row_idx, end_column=col_idx)
                            except ValueError:
                                pass
                                
                        cell = ws.cell(row=row_idx, column=col_idx)
                        subj_label = rec.course_code
                        if rec.subject_type == SubjectType.TUTORIAL:
                            subj_label += " (T)"
                            
                        text = f"{batch}\n{subj_label}\n{' '.join(rec.room_nos) if rec.room_nos else '-'}\n{' '.join(rec.fac_ids) if rec.fac_ids else '-'}"
                        cell.value = text
                        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                        cell.fill = PRAC_FILL if rec.subject_type == SubjectType.PRACTICAL else TUT_FILL

        # Adjust column widths
        ws.column_dimensions['A'].width = 12
        for c in range(2, 2 + 5 * num_batches):
            ws.column_dimensions[get_column_letter(c)].width = 15

    # ── FACULTY SHEETS ───────────────────────────────────────
    # Create one combined sheet for all faculty or separate sheets?
    # Let's create one FACULTY_MASTER sheet to avoid too many sheets
    ws_fac = wb.create_sheet(title="FACULTY MASTER")
    
    current_row = 1
    for fac_id, fac in data.faculty.items():
        week_records = timetable_index.get_faculty_week(fac_id)
        if not week_records: continue
        
        # Faculty header
        ws_fac.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=6)
        cell = ws_fac.cell(row=current_row, column=1, value=f"FACULTY: {fac_id}")
        cell.font = Font(bold=True, size=12)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        current_row += 1
        
        # Days header
        ws_fac.cell(row=current_row, column=1, value="Time").border = thin_border
        for i, day in enumerate(DAYS):
            ws_fac.cell(row=current_row, column=2+i, value=day).border = thin_border
        current_row += 1
        
        # Grid
        start_grid_row = current_row
        for p_idx, p_label in enumerate(PERIODS):
            r = start_grid_row + p_idx
            ws_fac.cell(row=r, column=1, value=p_label).border = thin_border
            for c in range(2, 7):
                ws_fac.cell(row=r, column=c).border = thin_border
                
            if p_idx == 4: # Lunch
                ws_fac.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
                lc = ws_fac.cell(row=r, column=2, value="LUNCH")
                lc.alignment = Alignment(horizontal="center")
                lc.fill = LUNCH_FILL
                
        # Fill records
        for slot_id, rec_list in week_records.items():
            day_idx = slot_id // 9
            period = slot_id % 9 + 1
            if period == 5: continue
            
            for rec in rec_list:
                if slot_id != rec.slot_id: continue
                r = start_grid_row + period - 1
                c = 2 + day_idx
                
                if rec.duration > 1:
                    try:
                        ws_fac.merge_cells(start_row=r, start_column=c, end_row=r+rec.duration-1, end_column=c)
                    except ValueError:
                        pass
                
                cell = ws_fac.cell(row=r, column=c)
                text = f"{rec.course_code}\n{rec.division_key}\n{' '.join(rec.room_nos)}"
                cell.value = text
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                if rec.subject_type == SubjectType.THEORY: cell.fill = THEORY_FILL
                elif rec.subject_type == SubjectType.PRACTICAL: cell.fill = PRAC_FILL
                else: cell.fill = TUT_FILL
                
        current_row = start_grid_row + 9 + 2 # Add spacing

    ws_fac.column_dimensions['A'].width = 12
    for c in range(2, 7):
        ws_fac.column_dimensions[get_column_letter(c)].width = 15

    # ── ROOM SHEETS ──────────────────────────────────────────
    ws_room = wb.create_sheet(title="ROOM MASTER")
    
    current_row = 1
    for room_no, room in data.rooms.items():
        if not room_no or room_no in ("-", "ONLINE", "CR"):
            continue
        week_records = timetable_index.get_room_week(room_no)
        if not week_records: continue
        
        # Room header
        ws_room.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=6)
        cell = ws_room.cell(row=current_row, column=1, value=f"ROOM: {room_no} ({room.room_type})")
        cell.font = Font(bold=True, size=12)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        current_row += 1
        
        # Days header
        ws_room.cell(row=current_row, column=1, value="Time").border = thin_border
        for i, day in enumerate(DAYS):
            ws_room.cell(row=current_row, column=2+i, value=day).border = thin_border
        current_row += 1
        
        # Grid
        start_grid_row = current_row
        for p_idx, p_label in enumerate(PERIODS):
            r = start_grid_row + p_idx
            ws_room.cell(row=r, column=1, value=p_label).border = thin_border
            for c in range(2, 7):
                ws_room.cell(row=r, column=c).border = thin_border
                
            if p_idx == 4: # Lunch
                ws_room.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
                lc = ws_room.cell(row=r, column=2, value="LUNCH")
                lc.alignment = Alignment(horizontal="center")
                lc.fill = LUNCH_FILL
                
        # Fill records
        for slot_id, rec_list in week_records.items():
            day_idx = slot_id // 9
            period = slot_id % 9 + 1
            if period == 5: continue
            
            for rec in rec_list:
                if slot_id != rec.slot_id: continue
                r = start_grid_row + period - 1
                c = 2 + day_idx
                
                if rec.duration > 1:
                    try:
                        ws_room.merge_cells(start_row=r, start_column=c, end_row=r+rec.duration-1, end_column=c)
                    except ValueError:
                        pass
                
                cell = ws_room.cell(row=r, column=c)
                text = f"{rec.course_code}\n{rec.division_key}\n{' '.join(rec.fac_ids)}"
                cell.value = text
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                if rec.subject_type == SubjectType.THEORY: cell.fill = THEORY_FILL
                elif rec.subject_type == SubjectType.PRACTICAL: cell.fill = PRAC_FILL
                else: cell.fill = TUT_FILL
                
        current_row = start_grid_row + 9 + 2 # Add spacing

    ws_room.column_dimensions['A'].width = 12
    for c in range(2, 7):
        ws_room.column_dimensions[get_column_letter(c)].width = 15

    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()
