"""
FastAPI entrypoint — Multi-Department Timetable System

Original routes (preserved):
  GET  /          -> upload form (template-based)
  POST /generate  -> run scheduler, return HTML timetable grid

JSON API — Department (requires dept Bearer token):
  POST /api/auth/login
  POST /api/auth/logout
  GET  /api/me
  POST /api/dept/upload
  POST /api/dept/solve
  GET  /api/dept/timetable
  GET  /api/dept/timetable/faculty/{fac_id}
  GET  /api/dept/timetable/room/{room_no}
  GET  /api/dept/conflicts
  GET  /api/dept/export

JSON API — Admin (requires admin Bearer token):
  GET  /api/admin/timetables
  GET  /api/admin/timetable/{id}
  PUT  /api/admin/timetable/{id}/remark
  POST /api/admin/timetable/{id}/regenerate
  GET  /api/admin/export/{id}
  GET  /api/admin/logs
  GET  /api/admin/departments

Legacy single-state JSON API (kept for backward compat):
  POST /api/upload
  POST /api/solve
  GET  /api/timetable
  GET  /api/timetable/faculty/{fac_id}
  GET  /api/timetable/room/{room_no}
  GET  /api/conflicts
  GET  /api/export
"""

import os
import io
import time as _time
import tempfile
from typing import Optional

from fastapi import FastAPI, UploadFile, File, HTTPException, Request
from fastapi.responses import HTMLResponse, FileResponse, JSONResponse, Response
from fastapi.templating import Jinja2Templates
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

from app.loader import (
    load_subjects, load_classrooms, load_fixed_slots, load_faculty_unavailability,
    extract_division_meta, build_practical_sessions, build_tutorial_sessions
)
from app.scheduler import run_full_schedule
from app.models import Timetable, ScheduleEntry, VALID_DAYS, VALID_PERIODS, PERIOD_TIME_MAP
from app import db, auth

app = FastAPI(title="Multi-Dept Timetable Scheduler")
templates = Jinja2Templates(directory="templates")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ------------------------------------------------------------------
# Initialise DB on startup
# ------------------------------------------------------------------

@app.on_event("startup")
def on_startup():
    db.init_db()


# ------------------------------------------------------------------
# Per-department in-memory state  {dept_id -> state_dict}
# ------------------------------------------------------------------

def _empty_state() -> dict:
    return {
        "subjects": [],
        "classrooms": [],
        "fixed_slots": [],
        "faculty_unavailability": [],
        "timetable": None,
        "unscheduled": [],
        "bottlenecks": [],
        "division_meta": None,
        "practical_sessions": None,
        "tutorial_sessions": None,
        "solve_time_ms": 0,
        "tmp_path": None,
        "last_tt_id": None,   # DB id of the last saved timetable
    }

_dept_state: dict[str, dict] = {}  # dept_id -> state


def _get_dept_state(dept_id: str) -> dict:
    if dept_id not in _dept_state:
        _dept_state[dept_id] = _empty_state()
    return _dept_state[dept_id]


# Legacy single-state (backward compat)
_state: dict = _empty_state()


# ------------------------------------------------------------------
# Auth helpers
# ------------------------------------------------------------------

def _require_auth(request: Request):
    """Returns (dept_row, token). Raises 401 if not authenticated."""
    token = auth.get_token_from_request(request)
    if not token:
        raise HTTPException(status_code=401, detail="Authentication required")
    dept_id = auth.verify_token(token)
    if not dept_id:
        raise HTTPException(status_code=401, detail="Invalid or expired token")
    row = db.get_dept(dept_id)
    if not row:
        raise HTTPException(status_code=401, detail="Department not found")
    return dict(row), token


def _require_admin(request: Request):
    """Like _require_auth but also checks is_admin."""
    dept_info, token = _require_auth(request)
    if not dept_info.get("is_admin"):
        raise HTTPException(status_code=403, detail="Admin access required")
    return dept_info, token


def _get_client_ip(request: Request) -> str:
    forwarded = request.headers.get("X-Forwarded-For")
    if forwarded:
        return forwarded.split(",")[0].strip()
    return request.client.host if request.client else ""


# ------------------------------------------------------------------
# Shared grid/timetable helpers
# ------------------------------------------------------------------

DAY_INDEX = {d: i for i, d in enumerate(VALID_DAYS)}


def _entry_to_dict(entry: ScheduleEntry) -> dict:
    subj = entry.subject
    day_idx = DAY_INDEX.get(entry.timeslot.day, 0)
    period = entry.timeslot.start_period
    if period <= 4:
        frontend_period = period
    else:
        frontend_period = period + 1  # skip lunch row

    slot_id = day_idx * 9 + (frontend_period - 1)

    stype = "THEORY"
    if subj.subject_type == "P":
        stype = "PRACTICAL"
    elif subj.subject_type == "T":
        stype = "TUTORIAL"

    return {
        "slot_id": slot_id,
        "course_code": subj.name,
        "subject_type": stype,
        "faculty": [entry.faculty_code] if entry.faculty_code else [],
        "rooms": [entry.classroom.room_id],
        "batches": [entry.batch] if entry.batch != "ALL" else ["ALL"],
        "duration": subj.duration,
        "division": f"{subj.year}-{subj.branch}-{subj.division}",
    }


def _build_timetable_json(timetable: Timetable) -> dict:
    division_map: dict[str, list[dict]] = {}
    all_faculty: set[str] = set()
    all_rooms: set[str] = set()

    year_div_to_branches = {}
    for entry in timetable.entries:
        y, b, d = entry.subject.class_key
        if b != "ALL":
            year_div_to_branches.setdefault((y, d), set()).add(b)

    for entry in timetable.entries:
        d = _entry_to_dict(entry)
        y, b, div = entry.subject.class_key

        if b == "ALL":
            targets = [f"{y}-{br}-{div}" for br in year_div_to_branches.get((y, div), [])]
            if not targets:
                targets = [d["division"]]
        else:
            targets = [d["division"]]

        for div_key in targets:
            d_copy = d.copy()
            d_copy["division"] = div_key
            division_map.setdefault(div_key, []).append(d_copy)

        if entry.faculty_code:
            all_faculty.add(entry.faculty_code)
        all_rooms.add(entry.classroom.room_id)

    division_list = sorted(division_map.keys())
    divisions = {k: {"grid": division_map[k]} for k in division_list}

    return {
        "division_list": division_list,
        "faculty_list": sorted(all_faculty),
        "room_list": sorted(all_rooms),
        "divisions": divisions,
    }


def _build_export_excel(timetable: Timetable) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    wb.remove(wb.active)

    # Grid data
    grid: dict = {}
    year_div_to_branches = {}
    for entry in timetable.entries:
        y, b, d = entry.subject.class_key
        if b != "ALL":
            year_div_to_branches.setdefault((y, d), set()).add(b)

    for entry in timetable.entries:
        y, b, d = entry.subject.class_key
        if b == "ALL":
            targets = [(y, br, d) for br in year_div_to_branches.get((y, d), [])]
            if not targets:
                targets = [entry.subject.class_key]
        else:
            targets = [entry.subject.class_key]

        for key in targets:
            if key not in grid:
                grid[key] = {}
            for period in entry.timeslot.occupied_periods:
                slot_key = (entry.timeslot.day, period)
                grid[key].setdefault(slot_key, []).append(entry)

    fac_grid: dict = {}
    for entry in timetable.entries:
        fac = entry.faculty_code
        if not fac:
            continue
        day_map = fac_grid.setdefault(fac, {})
        for p in entry.timeslot.occupied_periods:
            day_map.setdefault((entry.timeslot.day, p), []).append(entry)

    room_grid: dict = {}
    for entry in timetable.entries:
        rm = entry.classroom.room_id
        if not rm:
            continue
        day_map = room_grid.setdefault(rm, {})
        for p in entry.timeslot.occupied_periods:
            day_map.setdefault((entry.timeslot.day, p), []).append(entry)

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    period_labels = []
    for p in sorted(PERIOD_TIME_MAP.keys()):
        t = PERIOD_TIME_MAP[p]
        period_labels.append((p, f"P{p} ({t[0].strftime('%H:%M')}-{t[1].strftime('%H:%M')})"))

    def populate_ws(ws, day_period_map, formatter):
        ws.cell(row=1, column=1, value="Period").font = header_font
        ws.cell(row=1, column=1).fill = header_fill
        ws.cell(row=1, column=1).alignment = center_align
        ws.cell(row=1, column=1).border = thin_border

        for di, day in enumerate(VALID_DAYS):
            cell = ws.cell(row=1, column=di + 2, value=day)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        for ri, (period, label) in enumerate(period_labels):
            row = ri + 2
            ws.cell(row=row, column=1, value=label).alignment = center_align
            ws.cell(row=row, column=1).border = thin_border
            ws.cell(row=row, column=1).font = Font(bold=True)

            for di, day in enumerate(VALID_DAYS):
                slot_key = (day, period)
                entries_at_slot = day_period_map.get(slot_key, [])
                cell_value = "\n---\n".join(formatter(e) for e in entries_at_slot) if entries_at_slot else ""
                cell = ws.cell(row=row, column=di + 2, value=cell_value)
                cell.alignment = center_align
                cell.border = thin_border

        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 22

    for class_key, day_period_map in grid.items():
        sheet_name = f"{class_key[0]}-{class_key[1]}-{class_key[2]}"[:31]
        ws = wb.create_sheet(title=sheet_name)

        def div_format(e, _ck=class_key):
            line = e.subject.name
            if e.batch != "ALL":
                line += f" (B{e.batch})"
            return f"{line}\n{e.faculty_code} | {e.classroom.room_id}"

        populate_ws(ws, day_period_map, div_format)

    for fac, day_period_map in fac_grid.items():
        sheet_name = f"F_{fac}"[:31].replace("/", "_").replace("\\", "_")
        ws = wb.create_sheet(title=sheet_name)

        def fac_format(e):
            c = e.subject.class_key
            div_str = f"{c[0]}-{c[1]}-{c[2]}"
            line = e.subject.name
            if e.batch != "ALL":
                line += f" (B{e.batch})"
            return f"{line}\n{div_str} | {e.classroom.room_id}"

        populate_ws(ws, day_period_map, fac_format)

    for rm, day_period_map in room_grid.items():
        sheet_name = f"R_{rm}"[:31].replace("/", "_").replace("\\", "_")
        ws = wb.create_sheet(title=sheet_name)

        def rm_format(e):
            c = e.subject.class_key
            div_str = f"{c[0]}-{c[1]}-{c[2]}"
            line = e.subject.name
            if e.batch != "ALL":
                line += f" (B{e.batch})"
            return f"{line}\n{div_str} | {e.faculty_code}"

        populate_ws(ws, day_period_map, rm_format)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _run_solve_for_state(state: dict) -> tuple[str, int, bytes]:
    """Run scheduler on given state. Returns (status, elapsed_ms, excel_bytes)."""
    t0 = _time.time()
    timetable, unscheduled, bottlenecks = run_full_schedule(
        state["subjects"],
        state["classrooms"],
        state["fixed_slots"],
        state["faculty_unavailability"],
        state.get("division_meta"),
        state.get("practical_sessions"),
        state.get("tutorial_sessions"),
    )
    elapsed_ms = int((_time.time() - t0) * 1000)

    state["timetable"] = timetable
    state["unscheduled"] = unscheduled
    state["bottlenecks"] = bottlenecks
    state["solve_time_ms"] = elapsed_ms

    status = "OPTIMAL" if len(unscheduled) == 0 else "FEASIBLE"
    excel_bytes = _build_export_excel(timetable)
    return status, elapsed_ms, excel_bytes


def _load_from_xlsx(path: str, state: dict) -> dict:
    """Parse xlsx into state. Returns stats dict."""
    subjects = load_subjects(path)
    classrooms = load_classrooms(path)
    fixed_slots = []
    faculty_unavailability = []

    try:
        fixed_slots = load_fixed_slots(path)
    except Exception as e:
        print(f"[WARN] Fixed slots not loaded: {e}")

    try:
        faculty_unavailability = load_faculty_unavailability(path)
    except Exception as e:
        print(f"[WARN] Faculty unavailability not loaded: {e}")

    division_meta = extract_division_meta(subjects)
    practical_sessions = build_practical_sessions(subjects, division_meta)
    tutorial_sessions = build_tutorial_sessions(subjects, division_meta)

    state["subjects"] = subjects
    state["classrooms"] = classrooms
    state["fixed_slots"] = fixed_slots
    state["faculty_unavailability"] = faculty_unavailability
    state["division_meta"] = division_meta
    state["practical_sessions"] = practical_sessions
    state["tutorial_sessions"] = tutorial_sessions
    state["timetable"] = None
    state["unscheduled"] = []
    state["bottlenecks"] = []

    divisions = set()
    courses = set()
    faculty_set = set()
    for s in subjects:
        divisions.add(f"{s.year}-{s.branch}-{s.division}")
        courses.add(s.name)
        for f in s.faculty_codes:
            faculty_set.add(f)

    return {
        "divisions": len(divisions),
        "courses": len(courses),
        "faculty": len(faculty_set),
        "rooms": len(set(c.room_id for c in classrooms)),
    }


# ====================================================================
# Auth routes
# ====================================================================

class LoginBody(BaseModel):
    dept_id: str
    password: str


@app.post("/api/auth/login")
async def api_login(body: LoginBody, request: Request):
    if not db.verify_password(body.dept_id, body.password):
        raise HTTPException(status_code=401, detail="Invalid credentials")

    token = auth.create_token(body.dept_id)
    row = db.get_dept(body.dept_id)
    db.log_action(body.dept_id, "login", _get_client_ip(request))

    return JSONResponse({
        "token": token,
        "dept_id": body.dept_id,
        "display_name": row["display_name"],
        "is_admin": bool(row["is_admin"]),
    })


@app.post("/api/auth/logout")
async def api_logout(request: Request):
    dept_info, token = _require_auth(request)
    db.log_action(dept_info["dept_id"], "logout", _get_client_ip(request))
    auth.revoke_token(token)
    return JSONResponse({"success": True})


@app.get("/api/me")
async def api_me(request: Request):
    dept_info, _ = _require_auth(request)
    return JSONResponse({
        "dept_id": dept_info["dept_id"],
        "display_name": dept_info["display_name"],
        "is_admin": bool(dept_info["is_admin"]),
    })


# ====================================================================
# Department routes  (/api/dept/...)
# ====================================================================

@app.post("/api/dept/upload")
async def dept_upload(request: Request, file: UploadFile = File(...)):
    dept_info, _ = _require_auth(request)
    dept_id = dept_info["dept_id"]

    if dept_info.get("is_admin"):
        raise HTTPException(status_code=403, detail="Admins don't upload; use a department account")
    if not file.filename.endswith((".xlsx", ".xls")):
        return JSONResponse({"success": False, "error": "Please upload an Excel file (.xlsx)"})

    state = _get_dept_state(dept_id)

    # Clean old temp
    if state["tmp_path"] and os.path.exists(state["tmp_path"]):
        try:
            os.unlink(state["tmp_path"])
        except OSError:
            pass

    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        tmp.write(await file.read())
        tmp_path = tmp.name

    state["tmp_path"] = tmp_path

    try:
        stats = _load_from_xlsx(tmp_path, state)
        db.log_action(dept_id, f"upload:{file.filename}", _get_client_ip(request))
        return JSONResponse({"success": True, "data": stats})
    except Exception as e:
        return JSONResponse({"success": False, "error": str(e)})


@app.post("/api/dept/solve")
async def dept_solve(request: Request):
    dept_info, _ = _require_auth(request)
    dept_id = dept_info["dept_id"]

    if dept_info.get("is_admin"):
        raise HTTPException(status_code=403, detail="Use admin regenerate endpoint")

    state = _get_dept_state(dept_id)
    if not state["subjects"]:
        raise HTTPException(status_code=400, detail="No data uploaded yet")

    try:
        status, elapsed_ms, excel_bytes = _run_solve_for_state(state)
        tt_id = db.save_timetable(dept_id, status, excel_bytes)
        state["last_tt_id"] = tt_id
        db.log_action(dept_id, f"generate:status={status}", _get_client_ip(request))

        return JSONResponse({
            "status": status,
            "solve_time_ms": elapsed_ms,
            "timetable_id": tt_id,
            "total_entries": len(state["timetable"].entries),
            "unscheduled_count": len(state["unscheduled"]),
        })
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/api/dept/timetable")
async def dept_timetable(request: Request):
    dept_info, _ = _require_auth(request)
    dept_id = dept_info["dept_id"]
    state = _get_dept_state(dept_id)

    if state["timetable"] is None:
        raise HTTPException(status_code=400, detail="Timetable not generated yet")

    return JSONResponse(_build_timetable_json(state["timetable"]))


@app.get("/api/dept/timetable/faculty/{fac_id}")
async def dept_faculty_timetable(fac_id: str, request: Request):
    dept_info, _ = _require_auth(request)
    state = _get_dept_state(dept_info["dept_id"])
    if state["timetable"] is None:
        raise HTTPException(status_code=400, detail="Timetable not generated yet")
    entries = state["timetable"].get_by_faculty(fac_id)
    return JSONResponse({"faculty": fac_id, "grid": [_entry_to_dict(e) for e in entries]})


@app.get("/api/dept/timetable/room/{room_no}")
async def dept_room_timetable(room_no: str, request: Request):
    dept_info, _ = _require_auth(request)
    state = _get_dept_state(dept_info["dept_id"])
    if state["timetable"] is None:
        raise HTTPException(status_code=400, detail="Timetable not generated yet")
    entries = [e for e in state["timetable"].entries if e.classroom.room_id == room_no]
    return JSONResponse({"room": room_no, "grid": [_entry_to_dict(e) for e in entries]})


@app.get("/api/dept/conflicts")
async def dept_conflicts(request: Request):
    dept_info, _ = _require_auth(request)
    state = _get_dept_state(dept_info["dept_id"])
    conflicts: dict = {}

    if state["unscheduled"]:
        msgs = []
        for s in state["unscheduled"]:
            fac = ", ".join(s.faculty_codes) if s.faculty_codes else "N/A"
            msgs.append(f"{s.name} ({s.year} {s.branch} {s.division}, batch={s.batch}, type={s.subject_type}, faculty={fac})")
        conflicts["Unscheduled Subjects"] = msgs

    if state["bottlenecks"]:
        msgs = []
        for b in state["bottlenecks"]:
            msg = f"{b['division']} - {b['session']} could not find a slot."
            if b.get("failures"):
                msg += f" Reasons: {', '.join(f[1] for f in b['failures'])}"
            msgs.append(msg)
        conflicts["Session Bottlenecks"] = msgs

    return JSONResponse({"conflicts": conflicts})


@app.get("/api/dept/export")
async def dept_export(request: Request):
    dept_info, _ = _require_auth(request)
    dept_id = dept_info["dept_id"]
    state = _get_dept_state(dept_id)

    if state["last_tt_id"] is None:
        raise HTTPException(status_code=400, detail="No timetable saved yet")

    blob = db.get_timetable_blob(state["last_tt_id"])
    if not blob:
        raise HTTPException(status_code=404, detail="Export not found")

    db.log_action(dept_id, "export", _get_client_ip(request))
    return Response(
        content=blob,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={dept_id}_timetable.xlsx"},
    )


# ====================================================================
# Admin routes  (/api/admin/...)
# ====================================================================

@app.get("/api/admin/departments")
async def admin_departments(request: Request):
    _require_admin(request)
    rows = db.list_departments()
    return JSONResponse({"departments": [dict(r) for r in rows]})


@app.get("/api/admin/timetables")
async def admin_timetables(request: Request, dept_id: Optional[str] = None):
    _require_admin(request)
    rows = db.list_timetables(dept_id)
    return JSONResponse({"timetables": [dict(r) for r in rows]})


@app.get("/api/admin/timetable/{tt_id}")
async def admin_get_timetable(tt_id: int, request: Request):
    _require_admin(request)
    row = db.get_timetable(tt_id)
    if not row:
        raise HTTPException(status_code=404, detail="Timetable not found")

    # Try to get live in-memory timetable for grid view
    dept_id = row["dept_id"]
    state = _dept_state.get(dept_id)
    if state and state.get("timetable") and state.get("last_tt_id") == tt_id:
        grid_json = _build_timetable_json(state["timetable"])
    else:
        grid_json = None

    return JSONResponse({
        "id": row["id"],
        "dept_id": row["dept_id"],
        "generated_at": row["generated_at"],
        "status": row["status"],
        "admin_remark": row["admin_remark"],
        "grid": grid_json,
    })


class RemarkBody(BaseModel):
    remark: str


@app.put("/api/admin/timetable/{tt_id}/remark")
async def admin_set_remark(tt_id: int, body: RemarkBody, request: Request):
    dept_info, _ = _require_admin(request)
    row = db.get_timetable(tt_id)
    if not row:
        raise HTTPException(status_code=404, detail="Timetable not found")
    db.set_remark(tt_id, body.remark)
    db.log_action(dept_info["dept_id"], f"remark:tt_id={tt_id}", _get_client_ip(request))
    return JSONResponse({"success": True})


@app.post("/api/admin/timetable/{tt_id}/regenerate")
async def admin_regenerate(tt_id: int, request: Request):
    dept_info, _ = _require_admin(request)
    row = db.get_timetable(tt_id)
    if not row:
        raise HTTPException(status_code=404, detail="Timetable not found")

    dept_id = row["dept_id"]
    state = _dept_state.get(dept_id)
    if not state or not state.get("subjects"):
        raise HTTPException(
            status_code=400,
            detail=f"Department {dept_id} data not in memory. Ask the department to upload their file again."
        )

    try:
        status, elapsed_ms, excel_bytes = _run_solve_for_state(state)
        db.update_timetable_status(tt_id, status, excel_bytes)
        state["last_tt_id"] = tt_id
        db.log_action(dept_info["dept_id"], f"admin_regenerate:tt_id={tt_id}", _get_client_ip(request))

        return JSONResponse({
            "success": True,
            "status": status,
            "solve_time_ms": elapsed_ms,
            "total_entries": len(state["timetable"].entries),
        })
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/api/admin/export/{tt_id}")
async def admin_export(tt_id: int, request: Request):
    dept_info, _ = _require_admin(request)
    row = db.get_timetable(tt_id)
    if not row:
        raise HTTPException(status_code=404, detail="Timetable not found")

    blob = db.get_timetable_blob(tt_id)
    if not blob:
        raise HTTPException(status_code=404, detail="Excel not available for this timetable")

    db.log_action(dept_info["dept_id"], f"admin_export:tt_id={tt_id}", _get_client_ip(request))
    dept_id = row["dept_id"]
    return Response(
        content=blob,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={dept_id}_timetable_{tt_id}.xlsx"},
    )


@app.get("/api/admin/logs")
async def admin_logs(request: Request, limit: int = 200):
    _require_admin(request)
    rows = db.get_logs(limit)
    return JSONResponse({"logs": [dict(r) for r in rows]})


# ====================================================================
# Legacy single-state JSON API (backward compat)
# ====================================================================

@app.get("/", response_class=HTMLResponse)
async def upload_form(request: Request):
    return templates.TemplateResponse("upload.html", {"request": request})


@app.post("/generate", response_class=HTMLResponse)
async def generate_timetable(request: Request, file: UploadFile = File(...)):
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="Please upload an Excel file (.xlsx)")

    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        tmp.write(await file.read())
        tmp_path = tmp.name

    try:
        state = _empty_state()
        _load_from_xlsx(tmp_path, state)

        timetable, unscheduled, bottlenecks = run_full_schedule(
            state["subjects"], state["classrooms"], state["fixed_slots"],
            state["faculty_unavailability"], state["division_meta"],
            state["practical_sessions"], state["tutorial_sessions"]
        )

        grid_data: dict = {}
        year_div_to_branches = {}
        for entry in timetable.entries:
            y, b, d = entry.subject.class_key
            if b != "ALL":
                year_div_to_branches.setdefault((y, d), set()).add(b)

        for entry in timetable.entries:
            y, b, d = entry.subject.class_key
            if b == "ALL":
                targets = [(y, br, d) for br in year_div_to_branches.get((y, d), [])]
                if not targets:
                    targets = [entry.subject.class_key]
            else:
                targets = [entry.subject.class_key]

            for key in targets:
                if key not in grid_data:
                    grid_data[key] = {}
                for period in entry.timeslot.occupied_periods:
                    slot_key = (entry.timeslot.day, period)
                    grid_data[key].setdefault(slot_key, []).append(entry)

    finally:
        os.unlink(tmp_path)

    return templates.TemplateResponse("timetable.html", {
        "request": request,
        "grid_data": grid_data,
        "days": VALID_DAYS,
        "periods": list(PERIOD_TIME_MAP.keys()),
        "period_times": {p: f"{t[0].strftime('%H:%M')}-{t[1].strftime('%H:%M')}"
                         for p, t in PERIOD_TIME_MAP.items()},
        "unscheduled": unscheduled,
        "total": len(timetable.entries),
    })


@app.post("/api/upload")
async def api_upload(file: UploadFile = File(...)):
    if not file.filename.endswith((".xlsx", ".xls")):
        return JSONResponse({"success": False, "error": "Please upload an Excel file (.xlsx)"})

    if _state["tmp_path"] and os.path.exists(_state["tmp_path"]):
        try:
            os.unlink(_state["tmp_path"])
        except OSError:
            pass

    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        tmp.write(await file.read())
        tmp_path = tmp.name

    _state["tmp_path"] = tmp_path

    try:
        stats = _load_from_xlsx(tmp_path, _state)
        return JSONResponse({"success": True, "data": stats})
    except Exception as e:
        return JSONResponse({"success": False, "error": str(e)})


@app.post("/api/solve")
async def api_solve():
    if not _state["subjects"]:
        raise HTTPException(status_code=400, detail="No data uploaded yet. Upload an Excel file first.")

    status, elapsed_ms, _ = _run_solve_for_state(_state)

    return JSONResponse({
        "status": status,
        "solve_time_ms": elapsed_ms,
        "total_entries": len(_state["timetable"].entries),
        "unscheduled_count": len(_state["unscheduled"]),
    })


@app.get("/api/timetable")
async def api_timetable():
    if _state["timetable"] is None:
        raise HTTPException(status_code=400, detail="Timetable not generated yet.")
    return JSONResponse(_build_timetable_json(_state["timetable"]))


@app.get("/api/timetable/faculty/{fac_id}")
async def api_faculty_timetable(fac_id: str):
    if _state["timetable"] is None:
        raise HTTPException(status_code=400, detail="Timetable not generated yet.")
    entries = _state["timetable"].get_by_faculty(fac_id)
    return JSONResponse({"faculty": fac_id, "grid": [_entry_to_dict(e) for e in entries]})


@app.get("/api/timetable/room/{room_no}")
async def api_room_timetable(room_no: str):
    if _state["timetable"] is None:
        raise HTTPException(status_code=400, detail="Timetable not generated yet.")
    entries = [e for e in _state["timetable"].entries if e.classroom.room_id == room_no]
    return JSONResponse({"room": room_no, "grid": [_entry_to_dict(e) for e in entries]})


@app.get("/api/conflicts")
async def api_conflicts():
    unscheduled = _state.get("unscheduled", [])
    bottlenecks = _state.get("bottlenecks", [])
    conflicts: dict = {}

    if unscheduled:
        msgs = []
        for s in unscheduled:
            fac = ", ".join(s.faculty_codes) if s.faculty_codes else "N/A"
            msgs.append(f"{s.name} ({s.year} {s.branch} {s.division}, batch={s.batch}, type={s.subject_type}, faculty={fac})")
        conflicts["Unscheduled Subjects"] = msgs

    if bottlenecks:
        msgs = []
        for b in bottlenecks:
            msg = f"{b['division']} - {b['session']} could not find a slot."
            if b.get("failures"):
                msg += f" Reasons: {', '.join(f[1] for f in b['failures'])}"
            msgs.append(msg)
        conflicts["Session Bottlenecks"] = msgs

    return JSONResponse({"conflicts": conflicts})


@app.get("/api/export")
async def api_export():
    if _state["timetable"] is None:
        raise HTTPException(status_code=400, detail="Timetable not generated yet.")

    excel_bytes = _build_export_excel(_state["timetable"])
    export_path = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx").name
    with open(export_path, "wb") as f:
        f.write(excel_bytes)

    return FileResponse(
        export_path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename="timetable_export.xlsx",
    )


@app.get("/api/debug_subjects")
def debug_subjects():
    if not _state.get("subjects"):
        return {"error": "no state"}
    out = []
    for s in _state["subjects"]:
        if s.name in ["AIML", "CV", "MRAS", "AICL"]:
            out.append({
                "name": s.name, "branch": s.branch, "division": s.division,
                "course_type": s.course_type, "subject_type": s.subject_type,
                "room": s.room, "faculty_codes": s.faculty_codes, "hours_per_week": s.hours_per_week
            })
    return out
